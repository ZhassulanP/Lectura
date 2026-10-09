from io import BytesIO
from pathlib import Path
import os

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation as PowerPoint
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.ai.gemini_provider import get_provider
from app.core.database import get_db
from app.main import create_app
from app.models import Base
from app.schemas import CardsContent, NotesContent, QuizContent


def pdf_bytes(text="Energy is conserved. Entropy never decreases in an isolated system."):
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    if text:
        stream = DecodedStreamObject()
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream.set_data(f"BT /F1 12 Tf 50 700 Td ({escaped}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def pptx_bytes():
    deck = PowerPoint()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Energy"
    slide.placeholders[1].text = "Energy is conserved."
    deck.slides.add_slide(deck.slide_layouts[6])
    buffer = BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


class StubProvider:
    async def generate_notes(self, slides):
        return NotesContent.model_validate({"overview": "Energy conservation.", "sections": [{"heading": "Energy", "summary": "Energy is conserved.", "sourceSlides": [1]}]})

    async def generate_quiz(self, slides, options):
        return QuizContent.model_validate({"questions": [{
            "type": "MULTIPLE_CHOICE", "prompt": "What is conserved?", "options": ["Energy", "Entropy", "Heat", "Work"],
            "correctAnswer": "Energy", "explanation": "The lecture states energy is conserved.", "sourceSlides": [1],
        }, {
            "type": "SHORT_ANSWER", "prompt": "What never decreases in an isolated system?", "correctAnswer": "Entropy",
            "explanation": "The lecture states entropy never decreases.", "sourceSlides": [1],
        }][:options.question_count]})

    async def generate_flashcards(self, slides):
        return CardsContent.model_validate({"cards": [{"front": "Conserved quantity?", "back": "Energy", "sourceSlides": [1]}]})


@pytest.fixture
def client():
    url = os.environ.get("TEST_DATABASE_URL", "sqlite://")
    engine = create_engine(url, **({"connect_args": {"check_same_thread": False}, "poolclass": StaticPool} if url == "sqlite://" else {}))
    if url == "sqlite://":
        @event.listens_for(engine, "connect")
        def foreign_keys(connection, record):
            connection.execute("PRAGMA foreign_keys=ON")
    # TEST_DATABASE_URL must point to a dedicated disposable test database.
    Base.metadata.create_all(engine)
    app = create_app()
    def session():
        with Session(engine, expire_on_commit=False) as db:
            yield db
    app.dependency_overrides[get_db] = session
    app.dependency_overrides[get_provider] = StubProvider
    with TestClient(app) as test_client:
        yield test_client
    Base.metadata.drop_all(engine)
    engine.dispose()


@pytest.fixture
def presentation_id(client):
    response = client.post("/api/v1/presentations", files={"file": ("lecture.pdf", pdf_bytes(), "application/pdf")})
    assert response.status_code == 201, response.text
    return response.json()["id"]
